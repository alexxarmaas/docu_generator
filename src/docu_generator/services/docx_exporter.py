from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from docu_generator.models import GuideDraft, GuideSection


def _rgb(hex_value: str, fallback: str) -> RGBColor:
    value = (hex_value or fallback).lstrip("#")
    try:
        return RGBColor.from_string(value.upper())
    except Exception:
        return RGBColor.from_string(fallback.lstrip("#").upper())


def _section_blocks(section: GuideSection):
    if section.blocks:
        return section.blocks

    blocks = []
    from docu_generator.models import GuideBlock

    if section.body:
        blocks.append(GuideBlock(type="text", text=section.body))
    if section.checklist:
        blocks.append(GuideBlock(type="checklist", items=section.checklist))
    if section.note:
        blocks.append(
            GuideBlock(
                type="note",
                text=section.note,
                label="Importante",
                variant="warning",
            )
        )
    if section.image_name:
        blocks.append(
            GuideBlock(
                type="image",
                image_name=section.image_name,
                image_caption=section.image_caption,
            )
        )
    return blocks


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def _set_cell_border(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color.replace("#", ""))


def _add_callout(document: Document, label: str, text: str, variant: str, colors_map: dict):
    backgrounds = {
        "info": "EEF6FC",
        "tip": "F0FBFA",
        "warning": "FFF8E8",
        "success": "EEF9F1",
    }
    borders = {
        "info": "4A90C2",
        "tip": colors_map["accent"].replace("#", ""),
        "warning": "E2A72E",
        "success": "48A868",
    }

    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _shade_cell(cell, backgrounds.get(variant, backgrounds["info"]))
    _set_cell_border(cell, borders.get(variant, borders["info"]))

    p = cell.paragraphs[0]
    if label:
        run = p.add_run(label)
        run.bold = True
        run.font.color.rgb = colors_map["primary"]
        p.add_run("\n")

    p.add_run(text)
    document.add_paragraph()


def _image_width_inches(size: str) -> float:
    return {
        "small": 2.6,
        "medium": 4.2,
        "large": 5.8,
        "full": 6.6,
    }.get(size, 5.8)


def _alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(value, WD_ALIGN_PARAGRAPH.CENTER)


def build_docx(
    guide: GuideDraft,
    images: list[dict],
    profile: dict,
) -> bytes:
    document = Document()
    colors_map = {
        "primary": _rgb(
            profile.get("brand", {}).get("primary", "#092D54"),
            "#092D54",
        ),
        "accent": _rgb(
            profile.get("brand", {}).get("accent", "#00B8A9"),
            "#00B8A9",
        ),
    }

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    styles["Title"].font.color.rgb = colors_map["primary"]

    styles["Heading 1"].font.color.rgb = colors_map["primary"]
    styles["Heading 2"].font.color.rgb = colors_map["primary"]

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    image_lookup = {image["name"]: image for image in images}
    product = profile.get("product") or profile.get("name") or "Producto"

    if guide.show_cover:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.space_after = Pt(24)

        product_run = p.add_run(str(product).upper())
        product_run.bold = True
        product_run.font.size = Pt(12)
        product_run.font.color.rgb = colors_map["accent"]

        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.add_run(guide.title)

        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta.add_run(
            f"{guide.document_type}\nVersión {guide.version}\n{guide.status}"
        )
        meta_run.font.size = Pt(11)

        if guide.author:
            author = document.add_paragraph()
            author.alignment = WD_ALIGN_PARAGRAPH.CENTER
            author.add_run(f"Autor: {guide.author}")

        if guide.updated_at:
            updated = document.add_paragraph()
            updated.alignment = WD_ALIGN_PARAGRAPH.CENTER
            updated.add_run(f"Actualizado: {guide.updated_at}")

        document.add_page_break()

    if guide.show_toc and guide.sections:
        document.add_heading("Contenido", level=1)
        for index, step in enumerate(guide.sections, start=1):
            p = document.add_paragraph(style="List Number")
            p.add_run(step.title)
        document.add_page_break()

    document.add_heading(guide.title, level=1)

    if guide.introduction:
        document.add_paragraph(guide.introduction)

    for section_index, step in enumerate(guide.sections, start=1):
        document.add_heading(f"{section_index}. {step.title}", level=2)

        for block in _section_blocks(step):
            if block.type == "text":
                for paragraph in block.text.split("\n\n"):
                    if paragraph.strip():
                        document.add_paragraph(paragraph.strip())

            elif block.type == "checklist":
                for item in block.items:
                    p = document.add_paragraph()
                    p.style = document.styles["List Bullet"]
                    p.add_run(f"☐ {item}")

            elif block.type == "note":
                _add_callout(
                    document,
                    block.label or "Nota",
                    block.text,
                    block.variant,
                    colors_map,
                )

            elif block.type == "image" and block.image_name:
                image = image_lookup.get(block.image_name)
                if image:
                    p = document.add_paragraph()
                    p.alignment = _alignment(block.align)
                    run = p.add_run()
                    run.add_picture(
                        io.BytesIO(image["bytes"]),
                        width=Inches(_image_width_inches(block.image_width)),
                    )
                    if block.image_caption:
                        caption = document.add_paragraph()
                        caption.alignment = _alignment(block.align)
                        caption_run = caption.add_run(block.image_caption)
                        caption_run.italic = True
                        caption_run.font.size = Pt(8.5)

            elif block.type == "divider":
                p = document.add_paragraph()
                p_pr = p._p.get_or_add_pPr()
                p_bdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:color"), "DDE5EB")
                p_bdr.append(bottom)
                p_pr.append(p_bdr)

            elif block.type == "table":
                rows = [
                    [cell.strip() for cell in row.split("|")]
                    for row in block.items
                    if row.strip()
                ]
                if rows:
                    width = max(len(row) for row in rows)
                    table = document.add_table(rows=1, cols=width)
                    table.style = "Table Grid"
                    for col_index in range(width):
                        text = rows[0][col_index] if col_index < len(rows[0]) else ""
                        table.cell(0, col_index).text = text
                        _shade_cell(table.cell(0, col_index), "EEF3F6")

                    for source_row in rows[1:]:
                        cells = table.add_row().cells
                        for col_index in range(width):
                            cells[col_index].text = (
                                source_row[col_index]
                                if col_index < len(source_row)
                                else ""
                            )
                    document.add_paragraph()

            elif block.type == "pagebreak":
                document.add_page_break()

    if guide.closing_note:
        document.add_paragraph()
        _add_callout(
            document,
            "Finalización",
            guide.closing_note,
            "success",
            colors_map,
        )

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        f"{product} · {guide.document_type} · v{guide.version}"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(102, 122, 138)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
